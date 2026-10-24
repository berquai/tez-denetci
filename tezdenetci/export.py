"""Excel dışa aktarımı: Word'den yeniden üretilmiş dipnot listesi + Kaynak Kütüğü (hücre içi italikli)."""
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from .common import RText

FONT = 'Times New Roman'


def rich(rt):
    if rt is None or not rt.text:
        return None
    from .common import curly
    rt = curly(rt)
    if not any(rt.ital):
        return rt.text
    # yalnızca boşluktan oluşan parça Excel'de bozuk 'dize özelliği' sayılır: boşluklar soldaki karakterin biçimini alsın
    it = list(rt.ital)
    for k, ch in enumerate(rt.text):
        if ch.isspace():
            it[k] = it[k - 1] if k > 0 else next((it[j] for j in range(len(it)) if not rt.text[j].isspace()), False)
    rt = RText(rt.text, it)
    parts = []
    for s, it in rt.segments():
        parts.append(TextBlock(InlineFont(rFont=FONT, i=True, sz=11), s) if it else TextBlock(InlineFont(rFont=FONT, sz=11), s))
    return CellRichText(parts)


def masked(rt, mask):
    if mask is None:
        return rt
    return RText(rt.text, [bool(m) if m is not None else rt.ital[k] for k, m in enumerate(mask)])


def export_xlsx(A, path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Dipnotlar (Word)'
    head = ['', 'TAM DİPNOT', 'KISA DİPNOT', 'KAYNAKÇA', 'KAYNAK', 'DOSYA ADI', 'KİMLİK', 'BULGU', 'OLMASI GEREKEN']
    ws.append(head)
    by_fn = Counter(f.ref for f in A.findings if f.where == 'dipnot')
    for f in A.doc.footnotes:
        cits = A.fn_cits.get(f.num, [])
        fixed = A.fixed_footnote(f.num) if A.fn_ops.get(f.num) else None
        first = cits[0] if cits else None
        rows = []
        rt = f.rt.strip()
        if first is None:
            rows.append([f.num, rich(rt), None, None, None, None, None])
        else:
            for k, c in enumerate(cits):
                s = c.source
                is_first_use = A.src_first.get(s.sid) == f.num
                row = [f.num if k == 0 else None,
                       rich(rt) if (k == 0 and c.form == 'tam') else None,
                       rich(rt) if (k == 0 and c.form == 'kisa') else None,
                       rich(masked(s.kay, s.kay_mask)) if is_first_use and c.form == 'tam' else None,
                       s.link, s.dosya, s.sid]
                rows.append(row)
        for k, row in enumerate(rows):
            ws.append(row + [by_fn.get(f.num, 0) if k == 0 else None, rich(fixed.strip()) if (fixed and k == 0) else None])
            r = ws.max_row
            if row[4]:
                c = ws.cell(r, 5)
                c.value = 'KAYNAK GÖSTER'
                c.hyperlink = row[4]
                c.font = Font(color='0563C1', underline='single')
            if k == 0 and by_fn.get(f.num):
                ws.cell(r, 8).fill = PatternFill('solid', fgColor='FDE7E4')
    _fmt(ws, [6, 70, 50, 70, 15, 40, 9, 8, 70])

    ws2 = wb.create_sheet('Kaynak Kütüğü')
    ws2.append(['KİMLİK', 'TÜR', 'TAM (KÖK)', 'KISA (KÖK)', 'KAYNAKÇA', 'SAYFA ARALIĞI', 'DOSYA ADI', 'KAYNAK',
                'İLK DİPNOT (EXCEL)', 'KULLANIM', 'KISA BİÇİM SAYISI', 'WORD’DE İLK', 'NOT'])
    ex_notes = defaultdict(list)
    for f in A.excel_findings:
        if f.ref is not None:
            ex_notes[f.ref].append(f.msg)
    for s in A.sources:
        nums = [r.num for r in s.rows if isinstance(r.num, int)]
        notes = []
        for r in s.rows:
            notes.extend(ex_notes.get(r.row, []))
        ws2.append([s.sid, s.tip, rich(masked(s.tam, s.tam_mask)), rich(masked(s.kisa, s.kisa_mask)),
                    rich(masked(s.kay, s.kay_mask)),
                    f'{s.page_range[0]}-{s.page_range[1]}' if s.page_range else None, s.dosya, s.link,
                    min(nums) if nums else None, len(s.rows), len(s.kisa_variants),
                    A.src_first.get(s.sid), ' | '.join(dict.fromkeys(notes)) or None])
        r = ws2.max_row
        if s.link:
            c = ws2.cell(r, 8)
            c.value = 'KAYNAK GÖSTER'
            c.hyperlink = s.link
            c.font = Font(color='0563C1', underline='single')
    _fmt(ws2, [8, 9, 70, 50, 70, 10, 40, 15, 10, 9, 9, 9, 60])
    from . import AUTHOR, APP_NAME
    wb.properties.creator = AUTHOR
    wb.properties.title = 'Dipnot listesi ve Kaynak Kütüğü'
    wb.properties.description = f'{APP_NAME} ile üretildi'
    wb.save(path)


def _fmt(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, name=FONT)
        c.fill = PatternFill('solid', fgColor='EDE7DC')
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.font is None or not c.hyperlink:
                c.font = Font(name=FONT, size=11)
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'B2'
