"""Kaynakçayı Kaynak Kütüğü'nden (Excel kaynakça sütunu) yeniden üretir, düzeltir ve sıralar."""
import re
from dataclasses import dataclass, field
from rapidfuzz import fuzz
from .common import RText, norm, tr_sort_key, nfc
from .registry import expected_mask, top_split, GAZETTES
from .rules import caps_issues, title_regions, _q

VERIFIED_UNTIL = 500   # kullanıcı: Excel'in ilk 500 satırı gözle kontrol edildi


@dataclass
class BibEntry:
    source: object
    old: RText
    new: RText
    changes: list = field(default_factory=list)
    flags: list = field(default_factory=list)
    first_num: int = 0
    verified: bool = False
    key: tuple = ()


def _repl(rt, a, b, s):
    return rt.replace(a, b, RText(s, [rt.ital[a] if a < len(rt.ital) else False] * len(s)))


def normalize_text(rt, changes):
    t = rt
    if '\xa0' in t.text or '​' in t.text:
        n = t.text.count('\xa0') + t.text.count('​')
        t = RText(t.text.replace('\xa0', ' ').replace('​', ''),
                  [i for c, i in zip(t.text, t.ital) if c != '​'])
        changes.append(f'{n} görünmez karakter (bölünmez boşluk) temizlendi')
    m = re.search(r'(?<=\S) {2,}(?=\S)', t.text)
    while m:
        t = t.replace(m.start(), m.end(), RText(' ', [False]))
        if 'Çift boşluk giderildi' not in changes:
            changes.append('Çift boşluk giderildi')
        m = re.search(r'(?<=\S) {2,}(?=\S)', t.text)
    if '"' in t.text:
        txt = list(t.text)
        open_ = True
        for k, c in enumerate(txt):
            if c == '"':
                txt[k] = '“' if open_ else '”'
                open_ = not open_
        t = RText(''.join(txt), t.ital)
        changes.append('Düz tırnak (") → “ ”')
    if re.search(r"(?<=\w)'(?=\w)", t.text):
        t = RText(re.sub(r"(?<=\w)'(?=\w)", '’', t.text), t.ital)
        changes.append("Düz kesme (') → ’")
    for m in re.finditer(r'([,.;])”', t.text):
        pass
    if re.search(r'[,;]”', t.text):
        t = RText(re.sub(r'([,;])”', r'”\1', t.text), t.ital)
        changes.append('Tırnak içindeki virgül tırnak dışına alındı')
    st = t.text.rstrip()
    if st and not st.endswith('.'):
        t = t.slice(0, len(st)) + RText('.', [False])
        changes.append('Sona nokta eklendi')
    return t


def build_bibliography(A):
    entries = []
    for s in A.sources:
        if not s.kay.text.strip():
            continue
        nums = [r.num for r in s.rows if isinstance(r.num, int)]
        first = min(nums) if nums else 0
        changes, flags = [], []
        new = normalize_text(s.kay.strip(), changes)
        # italik: yalnızca kesin vakalar otomatik düzeltilir (tamamen düz kalmış dergi/kitap adı, Resmî Gazete)
        mask = expected_mask(new.text, s.tip, s.title_short)
        if mask is not None:
            it = list(new.ital)
            fixed = []
            for a0, b0 in _regions(mask, new.text):
                seg = new.text[a0:b0]
                if any(it[k] for k in range(a0, b0) if new.text[k].isalnum()):
                    continue
                from .registry import PUB_HINT
                gaz = any(g in seg for g in GAZETTES)
                decree = re.search(r'(Karar|Kararı|Kanun|Kanunu|Kararname|Kararnamesi|Yönetmelik|Yönetmeliği|Tebliğ|Tebliği)\b', seg)
                sure = gaz or (not re.search(r'\d', seg) and not PUB_HINT.search(seg) and not decree and (
                    s.tip == 'makale' or (s.tip == 'kitap' and s.title_short and norm(s.title_short) in norm(seg))))
                if not sure:
                    flags.append(f'{_q(seg, 60)} italik olmalı gibi görünüyor; kontrol et')
                    continue
                for k in range(a0, b0):
                    if mask[k] is not None:
                        it[k] = mask[k]
                    elif new.text[k] == ' ':
                        it[k] = True
                fixed.append(seg)
            if fixed:
                changes.append('İtalik yapıldı: ' + ', '.join(_q(x, 50) for x in fixed))
                new = RText(new.text, it)
        # büyük harf (gazete manşetleri muaf)
        regs = title_regions(new.text, mask, s.tip)
        if s.gazete:
            regs = [(a, b) for a, b in regs if not (new.text[a - 1:a] == '“')]
        iss = caps_issues(new.text, regs)
        if iss:
            t = new
            for a, b, w, sug in iss:
                t = RText(t.text[:a] + sug + t.text[b:], t.ital)
            changes.append('Büyük harf: ' + ', '.join(f'{w}→{sug}' for _, _, w, sug in iss[:8]) + ('…' if len(iss) > 8 else ''))
            new = t
        # makale sayfa aralığı
        if s.tip == 'makale' and not s.gazete and not re.search(r's\.\s*\d+\s*[-–]\s*\d+|Makale no|e\d{4,}', new.text) \
                and not re.search(r'http|Erişim', new.text):
            flags.append('Makale/bölüm girdisinde sayfa aralığı (s. X-Y) yok')
        # tez bilgileri
        if s.tez and not re.search(r'Enstitü', new.text):
            flags.append('Tez girdisinde enstitü / anabilim dalı / şehir eksik (danışman isteği)')
        from .common import curly
        new = curly(new)
        e = BibEntry(s, s.kay.strip(), new, changes, flags, first, 0 < first <= VERIFIED_UNTIL)
        e.key = bib_key(new.text)
        entries.append(e)
    entries.sort(key=lambda e: e.key)
    # olası mükerrerler
    for i, e in enumerate(entries):
        for f in entries[i + 1:i + 6]:
            na, nb = norm(e.new.text), norm(f.new.text)
            if re.sub(r'\d+', '', na) == re.sub(r'\d+', '', nb):
                continue   # yalnızca yıl/sayı farkı: farklı yıllık raporlar
            from difflib import SequenceMatcher
            ta, tb = na.split(), nb.split()
            diff = [w for op, i1, i2, j1, j2 in SequenceMatcher(None, ta, tb).get_opcodes() if op != 'equal'
                    for w in ta[i1:i2] + tb[j1:j2]]
            ORD = r'(\d+|[ivxlc]+|birinci|ikinci|üçüncü|dördüncü|beşinci|altıncı|yedinci|sekizinci|dokuzuncu|onuncu|cilt|kitap|c|s)'
            if diff and all(re.fullmatch(ORD, w) for w in diff):
                continue
            if set(re.findall(r'\d+', na)) != set(re.findall(r'\d+', nb)):
                continue
            if fuzz.ratio(na, nb) >= 95:
                e.flags.append(f'Olası mükerrer: {f.source.sid} ile çok benzer')
                f.flags.append(f'Olası mükerrer: {e.source.sid} ile çok benzer')
    return entries


def _regions(mask, text):
    out, k = [], 0
    while k < len(mask):
        if mask[k]:
            a = k
            while k < len(mask) and mask[k] is not False:
                k += 1
            b = k
            while b > a and not text[b - 1].isalnum():
                b -= 1
            if b > a:
                out.append((a, b))
        k += 1
    return out


def _runs(text, idxs):
    out, cur = [], []
    for k in sorted(idxs):
        if cur and k - cur[-1] > 2:
            out.append(text[cur[0]:cur[-1] + 1])
            cur = []
        cur.append(k)
    if cur:
        out.append(text[cur[0]:cur[-1] + 1])
    return out


def _n_authors(text):
    m = re.match(r'^((?:[^,“”()]+?, (?:[A-ZÇĞİÖŞÜ][a-zçğıöşü]?\.\s?)+(?:,\s*(?:ve\s+)?|\s+ve\s+))+)', text)
    if not m:
        return 1
    return max(1, len(re.findall(r'(?:[A-ZÇĞİÖŞÜ][a-zçğıöşü]?\.\s?)+(?:,|\s+ve)', m.group(1))))


def bib_key(text):
    """Türkçe abece; isimsiz (tırnakla başlayan) gazete haberleri başta; aynı ilk yazarda tek yazarlı eserler önce."""
    anon = text[:1] in '“"‘'
    first = text.lstrip('“"‘').split(', ')[0]
    rest = text.lstrip('“"‘')
    return (0 if anon else 1, tr_sort_key(first), _n_authors(text), tr_sort_key(rest))


def word_only_entries(A, entries):
    """Eski Word kaynakçasında olup kütükte karşılığı bulunmayan girdiler."""
    from rapidfuzz import process
    kn = [norm(e.new.text) for e in entries]
    out = []
    for b in A.doc.bib:
        bn = norm(b.rt.text)
        r = process.extractOne(bn, kn, scorer=fuzz.token_set_ratio, score_cutoff=85)
        if not r:
            out.append(b)
    return out


def export_docx(entries, template_path, out_path, style_id='akaynaka', heading='KAYNAKLAR'):
    """Kaynakçayı tezin kendi biçemiyle (şablon .docx'teki 'akaynaka' stili) Word dosyası olarak yazar."""
    import copy
    from docx import Document
    from docx.shared import Pt
    doc = Document(template_path)
    body = doc.element.body
    for el in list(body):
        if el.tag.endswith('}sectPr'):
            continue
        body.remove(el)
    style = None
    for st in doc.styles:
        if getattr(st, 'style_id', None) == style_id:
            style = st
            break
    head_style = None
    for st in doc.styles:
        if getattr(st, 'style_id', None) in ('asbezelSayfalarOrtal', 'Heading1'):
            head_style = st
            break
    h = doc.add_paragraph(heading, style=head_style) if head_style is not None else doc.add_paragraph(heading)
    for e in entries:
        p = doc.add_paragraph(style=style) if style is not None else doc.add_paragraph()
        for seg, it in e.new.segments():
            r = p.add_run(seg)
            r.italic = True if it else None
    from . import AUTHOR, APP_NAME
    doc.core_properties.author = AUTHOR
    doc.core_properties.title = 'Kaynakça'
    doc.core_properties.comments = f'{APP_NAME} ile Kaynak Kütüğü’nden üretildi'
    doc.save(out_path)
    _strip_footnotes(out_path)


def _strip_footnotes(path):
    """Şablondan gelen tez dipnotlarını pakettan temizler (yalnızca ayırıcılar kalır)."""
    import zipfile, shutil, tempfile, os
    from lxml import etree
    W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    tmp = tempfile.mktemp(suffix='.docx')
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in ('word/footnotes.xml', 'word/endnotes.xml'):
                root = etree.fromstring(data)
                for fn in list(root):
                    if fn.get(W + 'type') is None:
                        root.remove(fn)
                data = etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
            zout.writestr(item, data)
    shutil.move(tmp, path)
