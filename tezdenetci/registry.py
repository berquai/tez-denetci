"""Kaynak Kütüğü: Excel bilgiler listesinden yapılandırılmış kaynak kayıtları."""
import re
from dataclasses import dataclass, field
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from .common import RText, nfc, norm, tr_lower

# ---------------------------------------------------------------- Excel okuma

def cell_rt(cell):
    v = cell.value
    if v is None:
        return RText()
    base_it = bool(cell.font and cell.font.i)
    if isinstance(v, CellRichText):
        runs = []
        for part in v:
            if isinstance(part, TextBlock):
                it = part.font.i if part.font is not None and part.font.i is not None else base_it
                runs.append((part.text, bool(it)))
            else:
                runs.append((str(part), base_it))
        return RText.from_runs(runs)
    return RText.from_runs([(str(v), base_it)])


@dataclass
class ExRow:
    row: int
    num: object
    tam: RText
    kisa: RText
    kay: RText
    dosya: str
    link: str
    source: object = None


def load_excel(path, sheet=None):
    wb = openpyxl.load_workbook(path, rich_text=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.worksheets[0]
    rows = []
    for r in ws.iter_rows(min_row=2):
        if len(r) < 4:
            continue
        tam, kisa, kay = cell_rt(r[1]), cell_rt(r[2]), cell_rt(r[3])
        if not (tam.text.strip() or kisa.text.strip() or kay.text.strip()):
            continue
        link = None
        for c in r[4:6]:
            if c.hyperlink is not None:
                link = c.hyperlink.target
                break
        dosya = nfc(str(r[5].value).strip()) if len(r) > 5 and r[5].value else ''
        num = r[0].value
        rows.append(ExRow(r[0].row, num, tam.strip(), kisa.strip(), kay.strip(), dosya, link))
    return rows

# ---------------------------------------------------------------- atıf kalıbı ayrıştırma

LOC_RE = re.compile(r',\s*(?:s\.|ss\.|Madde|Geçici Madde|Paragraf|par\.|md\.)\s*[^,]*(?:,\s*(?:\d[\d\s\-–/]*|s\.\s*\d[^,]*))*\.?\s*$')
YEAR_RE = re.compile(r'\b(1[5-9]\d\d|20\d\d)\b')


def split_locator(rt):
    """'... , s. 293-295.' -> (kök RText, konum metni)"""
    t = rt.text.rstrip()
    m = LOC_RE.search(t)
    if m:
        stem = rt.slice(0, m.start())
        loc = t[m.start():].strip(' ,')
    else:
        stem = rt.slice(0, len(t))
        loc = ''
    # sondaki nokta
    while stem.text and stem.text[-1] in '. ':
        stem = stem.slice(0, len(stem) - 1)
    return stem, loc.rstrip('.')


GAZETTES = ('Resmî Ceride', 'Resmi Ceride', 'Resmî Gazete', 'Resmi Gazete', 'Düstur', 'Zabıt Ceridesi')
PUB_HINT = re.compile(
    r'(Yayın|Yayınları|Yayınevi|Kitabevi|Kitap\b|Ofset|WIDER|Kitapevi|Yayıncılık|Yayımları|Press|Publishing|Publisher|Publications|Office|'
    r'University|Üniversitesi|Matbaa|Basımevi|Vakfı|Bakanlığı|Müdürlüğü|Teşkilatı|Başkanlığı|Kurumu|Bank\b|Library|Ltd|Inc\b|Verlag|'
    r'[EÉ]ditions|Academic|Books|Routledge|Unwin|Brill|Blackwell|Wiley|Elgar|Springer|Palgrave|Macmillan|Sage|Pearson|Garden|Norton|'
    r'Yayımlanmamış|Doktora|Yüksek Lisans|Doctoral|Master|Dissertation|Thesis|Organization|Organisation|FAO\b|OECD|Bureau|Council|Commission|Institute|Enstitüsü|Series|Seri|Dizisi|Publication|Report No|Working Paper|Anabilim|Dalı|Bölümü|Fakültesi|Matbaası|Merkezi|Derneği|Odası|Birliği|T\.C\.|Department|Ministry|Network|Group|Foundation|Fund|Programme|CEPREMAP|NBER)',
    re.I)


def classify(stem_text):
    t = stem_text
    if 'http' in t or 'Erişim' in t:
        return 'web'
    if re.search(r'Kanun(?:u)? \(|Kanun No|Kanun Hükmünde|Kararname(?:si)? \(|Kararname No|Talimatname(?:si)? \(|Tüzük|Yönetmelik.*\(', t):
        return 'mevzuat'
    if re.search(r'Sayılı [^,]*?(Karar|Kanun|Kararname|Yönetmelik|Tebliğ|Genelge)|(Hakkında|Dair|İlişkin) (Karar|Kanun|Kararname|Yönetmelik|Tebliğ|Genelge)\b|Cumhurbaşkanlığı Kararnamesi|Bakanlar Kurulu Kararı', t) \
            and not re.match(r'^[^,]*,\s*[“"]', t):
        return 'mevzuat'
    if re.search(r'(^|,\s)[“"][^”"]+[”"](,\s|\s*\(|\s*$)', t):
        return 'makale'
    return 'kitap'


def _journal_span(text):
    """Kapanış tırnağından sonraki dergi/kitap/gazete adı aralığı."""
    m = re.search(r'[”"](?:,?\s*\([^)]*\))?,\s*', text)
    if not m:
        return None
    a = m.end()
    if text[a:a + 1] == '(':
        return None
    rest = text[a:]
    stop = re.search(r',\s*(?:[A-Z]{1,3}\s?\d|D\.|T\.|Dönem|Devre|Birleşim|İçtima|C\.|S\.|Cilt|Sayı|Vol\.|No\.|\(|Makale|Erişim|\d{1,2}[\s.]|(?:Ocak|Şubat|Mart|Nisan|Mayıs|Haziran|Temmuz|Ağustos|Eylül|Ekim|Kasım|Aralık|Birincikânun|İkincikânun|İlkkânun|Birinciteşrin|İkinciteşrin)\b|1[5-9]\d\d|20\d\d|https?:)', rest)
    pub = None
    for x, y in top_split(text, a):
        if x > a and (PUB_HINT.search(text[x:y]) and not text[x:y].startswith('(')):
            pub = x - 2
            break
    if stop:
        b = a + stop.start()
        if pub is not None and pub < b:
            b = pub
    elif pub is not None:
        b = pub
    else:
        b = len(text)
        while b > a and text[b - 1] in '. ':
            b -= 1
    if b - a < 2:
        return None
    return (a, b)


THESIS_HINT = re.compile(r'^(Yayımlanmamış|Yayınlanmamış|Doktora Tezi|Yüksek Lisans Tezi|Doctoral|Master|PhD|Ph\.D|Dissertation|Thesis|Unpublished)', re.I)


def top_split(text, start=0):
    """', ' ile böler ama parantez/tırnak içinde bölmez. [(chunk_start, chunk_end)]"""
    spans, depth, q, a = [], 0, False, start
    k = start
    while k < len(text):
        c = text[k]
        if c in '([':
            depth += 1
        elif c in ')]':
            depth = max(0, depth - 1)
        elif c == '“':
            q = True
        elif c == '”':
            q = False
        elif c == ',' and depth == 0 and not q and k + 1 < len(text) and text[k + 1].isspace():
            spans.append((a, k))
            a = k + 2
            k += 1
        k += 1
    spans.append((a, len(text)))
    return spans


def _post_title_chunk(ch):
    ch = ch.strip().rstrip('.')
    if not ch:
        return True
    if ch[0].islower() or ch.startswith(('and ', 've ', 'or ', 'veya ')):
        return False
    if ch.startswith('('):
        return True
    if re.fullmatch(r'(1[5-9]\d\d|20\d\d)( \(.*\))?(\.?\s*(DOI|doi|https?|Erişim).*)?', ch):
        return True
    if re.fullmatch(r'\d+(?:[-/.]\d+){2,}[\d\-/.\s]*', ch):   # arşiv yer numarası
        return True
    if re.fullmatch(r'[A-ZÇĞİÖŞÜ][a-zçğıöşüâîû]+', ch):   # şehir adı gibi tek kelime
        return 'city'
    if re.fullmatch(r'\d{1,2}[\s.]+\S+[\s.]+\d{4}', ch) or re.fullmatch(r'\S+ \d{4}', ch):
        return True
    if re.match(r'(s\.|ss\.|C\.|S\.|No\.|Sayı|Devre|Dönem|İçtima|Birleşim|Celse)', ch):
        return True
    return bool(PUB_HINT.search(ch)) or bool(THESIS_HINT.match(ch))


def _title_end(text, start, min_end):
    """Tırnaksız başlığın bitişi: kalan bütün parçaların yayın bilgisi olduğu ilk sınır."""
    chunks = top_split(text, 0)
    bounds = [a for a, b in chunks if a > min_end]
    for a in bounds:
        rest = [text[x:y] for x, y in chunks if x >= a]
        res = [_post_title_chunk(r) for r in rest]
        ok = all(res) and all(v != 'city' or (k == len(rest) - 2) for k, v in enumerate(res))
        if THESIS_HINT.match(rest[0].strip()) or ok:
            e = a - 2
            return e
    e = len(text)
    while e > start and text[e - 1] in '. ':
        e -= 1
    return e


@dataclass
class Source:
    sid: str
    tip: str
    tam: RText            # kök (sayfa bilgisi olmadan)
    kisa: RText
    kay: RText
    dosya: str
    link: str
    page_range: tuple = None
    tam_mask: list = None     # beklenen italik (karakter bazında, None = fark etmez)
    kisa_mask: list = None
    kay_mask: list = None
    title_short: str = ''
    kisa_variants: set = field(default_factory=set)
    rows: list = field(default_factory=list)
    keys: list = field(default_factory=list)   # (normal anahtar, biçim)
    notes: list = field(default_factory=list)
    gazete: bool = False
    tez: bool = False


MONTHS = r'(Ocak|Şubat|Mart|Nisan|Mayıs|Haziran|Temmuz|Ağustos|Eylül|Ekim|Kasım|Aralık|Birincikânun|İkincikânun|İlkkânun|Sonkânun|Birinciteşrin|İkinciteşrin|Teşrinievvel|Teşrinisani|Kânunuevvel|Kânunusani)'
GAZETE_RE = re.compile(r'[”"],\s*[^,“”]+,\s*(\d{1,2}\s+' + MONTHS + r'\s+\d{4}|\d{1,2}\.\d{1,2}\.\d{4})')
TEZ_RE = re.compile(r'(Yayımlanmamış|Yayınlanmamış)?\s*(Doktora|Yüksek Lisans|Sanatta Yeterlik)\s+Tezi')


def is_gazete(text):
    return bool(GAZETE_RE.search(text)) and not re.search(r'[”"],\s*[^,]+,\s*C\.', text)


def expected_mask(text, tip, title_short=''):
    """Bir kök metin için beklenen italik maskesi: True/False/None(fark etmez)."""
    n = len(text)
    mask = [False] * n
    spans = []
    if tip == 'mevzuat':
        for g in GAZETTES:
            for m in re.finditer(re.escape(g), text):
                spans.append((m.start(), m.end()))
    elif tip == 'makale':
        js = _journal_span(text)
        if js:
            spans.append(js)
    elif tip == 'kitap':
        if not title_short:
            return [None] * n
        if title_short:
            nt, idx = _norm_idx(text)
            ts = norm(title_short)
            k = nt.find(ts)
            if k >= 0 and ts:
                a = idx[k]
                b0 = idx[k + len(ts) - 1] + 1
                b = _title_end(text, a, b0)
                spans.append((a, b))
            else:
                return [None] * n
    elif tip == 'web':
        return [None] * n
    for a, b in spans:
        for k in range(a, b):
            mask[k] = True
    return _punct_rules(text, mask)


def _norm_idx(text):
    from .common import norm_map
    return norm_map(text)


def _punct_rules(text, mask):
    """Kurallar: parantez ve tırnak düz; italik bölgenin kenarındaki noktalama ve ekler düz;
    iç noktalama ve boşluk 'fark etmez'."""
    n = len(text)
    out = list(mask)
    for k, c in enumerate(text):
        if c in '()[]':
            # başlığın parçası olan parantez serbest (kullanıcı kararı)
            near = (k > 0 and mask[k - 1]) or (k + 1 < n and mask[k + 1])
            out[k] = None if near else False
        elif c in '“”"«»‘':
            # tırnak italikliği serbest (kullanıcı kararı)
            near = (k > 0 and mask[k - 1]) or (k + 1 < n and mask[k + 1])
            out[k] = None if near else False
        elif c == '’' or c == "'":
            # başlık içindeki kesme (Türkiye’de Kırsal ...) fark etmez; başlık sonundaki ek düz
            out[k] = None if (mask[k] and k + 1 < n and _inside_after_suffix(text, mask, k)) else False
        elif c.isspace():
            out[k] = None
        elif not c.isalnum():
            # kenarda mı içeride mi
            left = k > 0 and mask[k - 1]
            right = k + 1 < n and mask[k + 1]
            if mask[k] and left and right:
                out[k] = None
            elif c in ':-–/&.' and mask[k] and left:
                # başlık içi iki nokta, tire, kısaltma noktası
                j = k + 1
                while j < n and text[j].isspace():
                    j += 1
                out[k] = None if (j < n and mask[j]) else False
            else:
                out[k] = False
    return out


def _inside_after_suffix(text, mask, k):
    """Kesmeden sonra ek ve ardından italik devam ediyor mu?"""
    j = k + 1
    while j < len(text) and text[j].isalpha():
        j += 1
    while j < len(text) and text[j] == ' ':
        j += 1
    return j < len(text) and mask[j] if j < len(text) else False


def _short_title(kisa_stem, tam_stem):
    """Kısa kökten yazar kısmını ayırıp kısa başlığı bulur. Bulamazsa '' (bilinmiyor)."""
    parts = [p for p in re.split(r',\s+|\s+ve\s+(?=[A-ZÇĞİÖŞÜ][^\s,]+,)', kisa_stem) if p]
    tchunks = [tam_stem[a:b] for a, b in top_split(tam_stem)]
    i, t = 0, 0
    while i < len(parts):
        p = parts[i].strip()
        if p in ('vd.', 'vd', 'et al.', 've diğerleri'):
            i += 1
            break
        tc = norm(tchunks[t]) if t < len(tchunks) else ''
        np_ = norm(p)
        is_author = bool(np_) and tc.endswith(np_) and (len(tc) > len(np_) or i == 0)
        if is_author and i < len(parts) - 1:
            i += 1
            t += 1
        else:
            break
    nt = norm(tam_stem)
    for j in range(len(parts), i, -1):
        cand = ', '.join(parts[i:j])
        nc = norm(cand)
        if len(nc) >= 4 and not nc.isdigit() and nc in nt:
            return cand
    return ''


def build_registry(rows):
    sources = []
    by_kay = {}
    for r in rows:
        if r.kay.text.strip():
            key = norm(r.kay.text)
            if key in by_kay:
                src = by_kay[key]
                src.rows.append(r)
                r.source = src
                continue
            tam_stem, tam_loc = split_locator(r.tam) if r.tam.text else (RText(), '')
            kisa_stem, _ = split_locator(r.kisa) if r.kisa.text else (RText(), '')
            tip = classify(tam_stem.text or r.kay.text)
            pr = None
            m = re.search(r's\.\s*(\d+)\s*[-–]\s*(\d+)\.?\s*(?:DOI|$|\.)', r.kay.text)
            if m:
                pr = (int(m.group(1)), int(m.group(2)))
            src = Source(f'K{len(sources) + 1:04d}', tip, tam_stem, kisa_stem, r.kay, r.dosya, r.link, pr)
            src.title_short = _short_title(kisa_stem.text, tam_stem.text) if tip == 'kitap' else ''
            src.rows.append(r)
            r.source = src
            sources.append(src)
            by_kay[key] = src
    # kısa satırları kaynaklara bağla
    kisa_index = {}
    for s in sources:
        if s.kisa.text:
            kisa_index.setdefault(norm(s.kisa.text), []).append(s)
    for r in rows:
        if r.source is not None:
            continue
        stem, _ = split_locator(r.kisa if r.kisa.text else r.tam)
        cands = kisa_index.get(norm(stem.text), [])
        if not cands and r.tam.text:
            st2, _ = split_locator(r.tam)
            cands = [s for s in sources if s.tam.text and norm(s.tam.text) == norm(st2.text)]
        if len(cands) > 1 and r.dosya:
            c2 = [c for c in cands if c.dosya == r.dosya]
            cands = c2 or cands
        if not cands and r.dosya:
            same = [s_ for s_ in sources if s_.dosya == r.dosya]
            if len(same) == 1:
                cands = same
        if not cands:
            # dosya adıyla ve kısa kök benzerliğiyle son deneme
            from rapidfuzz import fuzz
            best, bs = None, 0
            for s in sources:
                if r.dosya and s.dosya != r.dosya:
                    continue
                sc = fuzz.ratio(norm(stem.text), norm(s.kisa.text))
                if sc > bs:
                    best, bs = s, sc
            if best is not None and bs >= 80:
                cands = [best]
        if cands:
            r.source = cands[0]
            cands[0].rows.append(r)
    for s in sources:
        s.gazete = s.tip == 'makale' and (is_gazete(s.tam.text) or is_gazete(s.kay.text))
        s.tez = bool(TEZ_RE.search(s.tam.text) or TEZ_RE.search(s.kay.text))
        s.tam_mask = expected_mask(s.tam.text, s.tip, s.title_short)
        s.kisa_mask = expected_mask(s.kisa.text, s.tip, s.title_short)
        s.kay_mask = expected_mask(s.kay.text, s.tip, s.title_short)
        for r in s.rows:
            if r.kisa.text:
                s.kisa_variants.add(split_locator(r.kisa)[0].text)
        keys = []
        if s.tam.text:
            keys.append((norm(s.tam.text), 'tam'))
        for v in s.kisa_variants | ({s.kisa.text} if s.kisa.text else set()):
            if v:
                keys.append((norm(v), 'kisa'))
        s.keys = sorted(set(keys), key=lambda x: -len(x[0]))
    return sources
